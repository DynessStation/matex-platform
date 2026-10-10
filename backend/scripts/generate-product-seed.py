"""Build the normalized MATEX product seed from the approved workbook.

Usage:
  python scripts/generate-product-seed.py <workbook.xlsx> [output.json]
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


CATEGORIES = {
    "LATEX MATTRESS": {
        "key": "latex-mattress",
        "id": "Kasur Lateks",
        "en": "Latex Mattress",
    },
    "BEDDING ACCESSORIES": {
        "key": "bedding-accessories",
        "id": "Aksesori Tempat Tidur",
        "en": "Bedding Accessories",
    },
    "THERAPY ACCESSORIES": {
        "key": "therapy-accessories",
        "id": "Aksesori Terapi",
        "en": "Therapy Accessories",
    },
    "OTOMOTIF": {
        "key": "automotive",
        "id": "Otomotif",
        "en": "Automotive",
    },
}


def text(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def number(value):
    if value in (None, ""):
        return None
    value = float(value)
    return int(value) if value.is_integer() else value


def slug(value):
    normalized = unicodedata.normalize("NFKD", text(value) or "")
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"(^-|-$)", "", re.sub(r"[^a-z0-9]+", "-", ascii_value))


def specs(value):
    raw = text(value)
    if not raw:
        return None
    candidate = raw if raw.startswith("{") else "{" + raw + "}"
    parsed = json.loads(candidate)
    if not isinstance(parsed, dict):
        raise ValueError("Specifications must resolve to an object")
    return parsed


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Workbook path is required")
    source = Path(sys.argv[1]).resolve()
    output = (
        Path(sys.argv[2]).resolve()
        if len(sys.argv) > 2
        else Path(__file__).resolve().parents[1]
        / "database"
        / "seeds"
        / "data"
        / "20261010_product_catalog.json"
    )

    sheet = load_workbook(source, data_only=True, read_only=True)["PRODUCT"]
    products = []
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=7, max_col=48, values_only=True), start=7
    ):
        values = list(row)
        if values[0] in (None, ""):
            continue
        category_source = text(values[1])
        if category_source not in CATEGORIES:
            raise ValueError(f"Unknown category at row {row_number}: {category_source}")
        sku = text(values[21])
        if not sku:
            raise ValueError(f"Missing SKU at row {row_number}")
        featured_source = number(values[37]) or 0
        products.append(
            {
                "source_row": row_number,
                "source_number": int(values[0]),
                "category_key": CATEGORIES[category_source]["key"],
                "key": slug(sku),
                "sku": sku,
                "product_type": "physical",
                "product_kind": text(values[22]),
                "unit": text(values[23]),
                "barcode": text(values[24]),
                "manufacturer_code": text(values[25]),
                "country_origin": text(values[26]),
                "origin_province": text(values[27]),
                "origin_city": text(values[28]),
                "hs_code": text(values[29]),
                "weight_grams": number(values[30]),
                "length_mm": number(values[31]),
                "width_mm": number(values[32]),
                "height_mm": number(values[33]),
                "min_order_qty": number(values[34]) or 1,
                "lead_time_days": number(values[35]),
                "stock_status": "preorder"
                if (text(values[36]) or "").lower().replace("-", " ") == "pre order"
                else "in_stock",
                "is_featured": bool(featured_source),
                "sort_order": int(values[0]),
                "status": "draft",
                "source_image_url": text(values[4]),
                "source_visibility": text(values[46]),
                "source_inputted_by": text(values[47]),
                "translations": {
                    "id-ID": {
                        "slug": slug(values[5]),
                        "name": text(values[2]),
                        "short_description": text(values[7]),
                        "description": text(values[9]),
                        "specifications": specs(values[11]),
                        "meta_title": text(values[13]),
                        "meta_description": text(values[15]),
                        "og_title": text(values[17]),
                        "og_description": text(values[19]),
                    },
                    "en-US": {
                        "slug": slug(values[6]),
                        "name": text(values[3]),
                        "short_description": text(values[8]),
                        "description": text(values[10]),
                        "specifications": specs(values[12]),
                        "meta_title": text(values[14]),
                        "meta_description": text(values[16]),
                        "og_title": text(values[18]),
                        "og_description": text(values[20]),
                    },
                },
                "prices": [
                    {"type": "end_user", "label": "Harga konsumen", "currency": "IDR", "amount": number(values[38]), "is_public": True},
                    {"type": "end_user", "label": "Consumer price", "currency": "USD", "amount": number(values[39]), "is_public": False},
                    {"type": "wholesale", "label": "Harga grosir", "currency": "IDR", "amount": number(values[40]), "is_public": False},
                    {"type": "wholesale", "label": "Wholesale price", "currency": "USD", "amount": number(values[41]), "is_public": False},
                ],
            }
        )

    if len(products) != 53:
        raise ValueError(f"Expected 53 products, found {len(products)}")
    sku_counts = Counter(item["sku"].lower() for item in products)
    if any(count > 1 for count in sku_counts.values()):
        raise ValueError("Duplicate SKU found")

    for locale in ("id-ID", "en-US"):
        counts = Counter(item["translations"][locale]["slug"] for item in products)
        for item in products:
            translation = item["translations"][locale]
            if counts[translation["slug"]] > 1:
                translation["slug"] = f'{translation["slug"]}-{slug(item["sku"])}'

    payload = {
        "schema_version": 1,
        "source": source.name,
        "source_date": "2026-10-10",
        "company_id": 1,
        "actor_admin_id": 1,
        "import_mode": "replace",
        "publication_status": "draft",
        "categories": [
            {
                "key": data["key"],
                "source_name": source_name,
                "sort_order": index,
                "translations": {
                    "id-ID": {"slug": data["key"], "name": data["id"]},
                    "en-US": {"slug": data["key"], "name": data["en"]},
                },
            }
            for index, (source_name, data) in enumerate(CATEGORIES.items(), 1)
        ],
        "products": products,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(products)} products and {len(payload['categories'])} categories to {output}")


if __name__ == "__main__":
    main()
